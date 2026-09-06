# -*- coding: utf-8 -*-
"""抽取最新版 xlsm 的分析2(机台产能) + 分析3(前十占比) + 更新日期，供看板内嵌"""
import openpyxl, json, os
from datetime import datetime, date

SRC = r"D:\WorkBuddy工作资料\injection_dashboard\最新排程产能分析-公式模板-最新.xlsm"
OUT = r"D:\WorkBuddy工作资料\injection_dashboard\analysis_data.json"

wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)

def num(v):
    if v is None: return None
    if isinstance(v,(int,float)): return v
    s=str(v).strip()
    if s in ("","#N/A","#VALUE!","None"): return None
    try: return float(s.replace(",",""))
    except: return None

def s(v):
    if v is None: return ""
    if isinstance(v,(datetime,date)): return v.strftime("%Y-%m-%d")
    return str(v).strip()

# ===== 分析2 机台产能 =====
ws2 = wb["分析2-机台产能分析"]
groups_def = [
    ("400T-550T", 3, 4, 5, 6),   # 实单, 6天, 7天, 预算 (1-based)
    ("230T-280T", 14, 15, 16, 17),
    ("150T",      24, 25, 26, 27),
    ("90T",       33, 34, 35, 36),
    ("150T-280T", 47, 48, 49, 50),
]
machine = {"weeks": [], "groups": {}}
for name, r, c6, c7, bud in groups_def:
    machine["groups"][name] = {"real": [], "cap6": [], "cap7": [], "budget": []}
for row in ws2.iter_rows(min_row=5, max_row=42, values_only=True):
    wk = s(row[1])
    if not wk.startswith("202"):
        continue
    machine["weeks"].append(wk)
    for name, r, c6, c7, bud in groups_def:
        g = machine["groups"][name]
        g["real"].append(num(row[r-1]))
        g["cap6"].append(num(row[c6-1]))
        g["cap7"].append(num(row[c7-1]))
        g["budget"].append(num(row[bud-1]))

# ===== 分析3 前十占比 =====
ws3 = wb["分析3-实单+预算待生产前十占比分析"]
shell = []   # 外壳前十: 品名/实单+预算PCS/金额/实单数量/预算数量
acc = []     # 配件前十: 品名/实单+预算PCS/金额
for r in ws3.iter_rows(min_row=17, max_row=40, values_only=True):
    sname = s(r[0])
    if sname == "" or sname.startswith("公式"):
        break
    shell.append({
        "name": sname,
        "pcs": num(r[1]),
        "amount": num(r[2]),
        "realQty": num(r[3]),
        "budgetQty": num(r[4]),
    })
    aname = s(r[5])
    if aname:
        acc.append({"name": aname, "pcs": num(r[6]), "amount": num(r[7]), "realQty": num(r[8]), "budgetQty": num(r[9])})

# ===== KPI 目标值（分析3 头部） =====
kpi = {}
for r in ws3.iter_rows(min_row=3, max_row=12, values_only=True):
    label = s(r[0])
    val = num(r[1])
    if label and val is not None:
        kpi[label] = val

# ===== 更新日期 =====
ws1 = wb["原始数据1（工单备料明细）"]
detailDate = ""
for r in ws1.iter_rows(min_row=1, max_row=1, values_only=True):
    for c in range(1, len(r)):
        if r[c-1] is not None and str(r[c-1]).startswith("更新日期") and c < len(r) and r[c] is not None:
            detailDate = s(r[c])
            break
# 原始数据1 更新日期（F1=更新日期, G1=日期）
if not detailDate:
    r1 = [c.value for c in ws1[1]]
    for i in range(len(r1)-1):
        if r1[i] is not None and "更新日期" in str(r1[i]):
            detailDate = s(r1[i+1])

result = {
    "machine": machine,
    "top10": {"shell": shell, "acc": acc},
    "kpi": kpi,
    "detailDate": detailDate,
}

os.makedirs(os.path.dirname(OUT), exist_ok=True)
with open(OUT, "w", encoding="utf-8") as f:
    json.dump(result, f, ensure_ascii=False)

print("机台周数:", len(machine["weeks"]), machine["weeks"][0], "~", machine["weeks"][-1])
for name, g in machine["groups"].items():
    print(f"  {name}: 实单样例={g['real'][:3]} 6天={g['cap6'][0]} 7天={g['cap7'][0]} 预算样例={g['budget'][:3]}")
print("外壳前十数:", len(shell), " 榜首:", shell[0] if shell else None)
print("配件前十数:", len(acc), " 榜首:", acc[0] if acc else None)
print("KPI:", kpi)
print("工单备料明细更新日期:", detailDate)
print("输出:", OUT, os.path.getsize(OUT), "bytes")
wb.close()
