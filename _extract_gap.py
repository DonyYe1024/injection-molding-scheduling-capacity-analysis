# -*- coding: utf-8 -*-
"""抽取 GAP 模具产能 + 原料需求 + 外壳库存 + 库存水位饼图 + 入库/年度汇总，输出 EMBEDDED_GAP 数据"""
import openpyxl, json, os, re
from datetime import datetime, date

SRC = r"D:\WorkBuddy工作资料\injection_dashboard\最新排程产能分析-公式模板-最新.xlsm"
OUT = r"D:\WorkBuddy工作资料\injection_dashboard\gap_material_data.json"

wbf = openpyxl.load_workbook(SRC, read_only=True, data_only=False)
wbv = openpyxl.load_workbook(SRC, read_only=True, data_only=True)

def s(v):
    if v is None: return ""
    if isinstance(v,(datetime,date)): return v.strftime("%Y-%m-%d")
    return str(v).strip()

def num(v):
    if v is None: return 0
    if isinstance(v,(int,float)): return v
    sv=str(v).strip()
    if sv in ("","#N/A","#VALUE!","None"): return 0
    try: return float(sv.replace(",",""))
    except: return 0

# GAP 表：值（read_only + iter_rows）
wsv = wbv["注塑模具产能问题GAP（王力预算）"]
gap_rows = list(wsv.iter_rows(values_only=True))

# 周表头（R4，D列 idx3 起）
hdr4 = [s(v) for v in gap_rows[3]]
weeks = []
for c in range(3, len(hdr4)):
    m = re.match(r'^(\d{4})WK(\d{1,2})', hdr4[c])
    if m: weeks.append(m.group(1)+'WK'+str(int(m.group(2))))
    elif weeks: break

workdays = [num(gap_rows[1][c]) for c in range(3, 3+len(weeks))]

# 公式（周期套数）—— data_only=False read_only
wsf = wbf["注塑模具产能问题GAP（王力预算）"]
gap_formula_rows = list(wsf.iter_rows(values_only=True))

molds = []
for i in range(4, 166):
    a = s(gap_rows[i][0])
    b = s(gap_rows[i][1])
    if not a or b != "实单": continue
    cap = ''
    for c in range(0, len(gap_formula_rows[i+3]) if i+3 < len(gap_formula_rows) else 0):
        v = gap_formula_rows[i+3][c]
        if isinstance(v, str) and v.startswith('=3600'):
            cap = v; break
    m = re.match(r'=3600/(\d+)\*22\*(\d+)', cap)
    cycle = int(m.group(1)) if m else 0
    sets = int(m.group(2)) if m else 0
    real = [num(gap_rows[i][c]) for c in range(3, 3+len(weeks))]
    budget = [num(gap_rows[i+1][c]) for c in range(3, 3+len(weeks))]
    molds.append({"mold": a.replace("\n"," ").strip(), "cycle": cycle, "sets": sets, "real": real, "budget": budget})

# 分析4（新表名）：原料需求 + 外壳库存水位 + 入库汇总 + 年度汇总
ws4v = wbv["分析4-原料需求-库存水位-入库产值统计"]
ws4f = wbf["分析4-原料需求-库存水位-入库产值统计"]
rows4v = list(ws4v.iter_rows(values_only=True))
rows4f = list(ws4f.iter_rows(values_only=True))

hdr4a = [s(v) for v in rows4v[3]]
mat_weeks = []
for c in range(5, 43):
    if c < len(hdr4a):
        m = re.match(r'^(\d{4})WK(\d{1,2})', hdr4a[c])
        if m: mat_weeks.append(m.group(1)+'WK'+str(int(m.group(2))))

materials = []
for i in range(4, 36):
    if i >= len(rows4v): break
    r = rows4v[i]
    code = s(r[1]); name = s(r[2])
    if not code or not name: break
    materials.append({
        "code": code, "name": name,
        "usage": s(r[0]), "desc": s(r[3]),
        "weekly": [num(r[c]) for c in range(5, 5+len(mat_weeks)) if c < len(r)],
    })

# 注塑外壳库存水位（分析4 D36:F42 表：D36:F36 表头 / D37:D41 五档 / D42 合计；E=数量PCS F=金额RMB G=占比）
stockWaterLevel = {
    "headers": [s(rows4v[35][3]), s(rows4v[35][4]), s(rows4v[35][5])] if len(rows4v) > 35 else ["", "", ""],
    "total": num(rows4v[41][4]) if len(rows4v) > 41 else 0,
    "totalAmount": num(rows4v[41][5]) if len(rows4v) > 41 else 0,
    "totalLabel": s(rows4v[41][3]) if len(rows4v) > 41 else "合计",
    "items": [],
}
for i in range(36, 41):  # Excel 37..41（0-based 36..40）
    if i >= len(rows4v): break
    r = rows4v[i]
    nm = s(r[3])
    if not nm: continue
    stockWaterLevel["items"].append({
        "name": nm,
        "qty": num(r[4]),
        "amount": num(r[5]),
        "ratio": num(r[6]),
    })

# 通用块抽取：defs=[(row_idx, 值起始列, 值结束列)]，读取 B列 label + 值区间；
# group 取 A列（分组头），达成率/目标KPI 行用自身 label 作 group。
def collect_block(rows, defs):
    out = []
    current_group = ""
    for ridx, c0, c1 in defs:
        if ridx >= len(rows): break
        r = rows[ridx]
        a = s(r[0])
        label = s(r[1])
        if not label: continue
        if a:
            current_group = a
        if "达成率" in label or "Rate" in label:
            group = "达成率(Rate)"
        elif "目标KPI" in label:
            group = "目标KPI"
        else:
            group = current_group
        values = [num(r[c]) for c in range(c0, c1)]
        out.append({"group": group, "label": label, "values": values})
    return out

# 入库汇总：D52~AG52 = 2026-09-01..09-30 每日日期；行 52..58（Excel53..59）
inbound = {
    "dates": [s(rows4v[51][c]) for c in range(3, 33)] if len(rows4v) > 51 else [],
    "rows": collect_block(rows4v, [
        (52, 3, 33), (53, 3, 33), (54, 3, 33),
        (55, 3, 33), (56, 3, 33), (57, 3, 33),
        (58, 3, 33),
    ]),
}

# 年度汇总：D62~O62 = 1月..12月；行 62..71（Excel63..72，跳过空行 65/66）
annual = {
    "months": [s(rows4v[61][c]) for c in range(3, 15)] if len(rows4v) > 61 else [],
    "rows": collect_block(rows4v, [
        (62, 3, 15), (63, 3, 15), (64, 3, 15),
        (67, 3, 15), (68, 3, 15), (69, 3, 15),
        (70, 3, 15), (71, 3, 15),
    ]),
}

# 原始数据9：外壳库存明细（旧 141 件来源已消失，改从原始数据9 抽取）
ws9v = wbv["原始数据9（注塑半成品原料库存水位）"]
rows9 = list(ws9v.iter_rows(values_only=True))
stock = []
seen = set()
stock_total = 0
for r in rows9[2:]:  # 跳过 更新日期 行 + 表头行
    if s(r[2]) != "塑胶外壳类": continue
    code = s(r[1])
    if not code: continue
    stock_total += num(r[9])          # J列 库存量
    if code in seen: continue
    seen.add(code)
    stock.append({"code": code, "name": s(r[2]), "spec": s(r[3])})

result = {
    "gap": {"weeks": weeks, "workdays": workdays, "molds": molds},
    "material": {"weeks": mat_weeks, "items": materials},
    "stock": stock,
    "stockTotal": stock_total,
    "stockWaterLevel": stockWaterLevel,
    "inbound": inbound,
    "annual": annual,
}
os.makedirs(os.path.dirname(OUT), exist_ok=True)
with open(OUT, "w", encoding="utf-8") as f:
    json.dump(result, f, ensure_ascii=False)

print("OK weeks=%d molds=%d materials=%d stock=%d stockTotal=%s stockWaterLevel=%d inbound_rows=%d annual_rows=%d bytes=%d" % (
    len(weeks), len(molds), len(materials), len(stock), stock_total,
    len(stockWaterLevel["items"]), len(inbound["rows"]), len(annual["rows"]), os.path.getsize(OUT)))
