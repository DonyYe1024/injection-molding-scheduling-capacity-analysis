# -*- coding: utf-8 -*-
"""从xlsm抽取聚合数据，输出JSON供HTML看板内嵌"""
import openpyxl, json, re
from datetime import datetime
from collections import defaultdict, Counter

SRC = r"D:\WorkBuddy工作资料\injection_dashboard\最新排程产能分析-公式模板-最新.xlsm"
OUT = r"D:\WorkBuddy工作资料\injection_dashboard\data.json"

wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)

def num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip()
    if s in ("", "#N/A", "#VALUE!", "#DIV/0!", "None"):
        return None
    try:
        return float(s.replace(",", ""))
    except:
        return None

# ============ 1. 机台产能（分析2） ============
ws2 = wb["分析2-机台产能分析"]
# 机台组定义: (名称, 实单列, 产能低列, 产能高列, 预算列, 产能标签)
groups = [
    ("400T-550T", "C", "D", "E", "F", "10台/11台×22h×6天"),
    ("230T-280T", "N", "O", "P", "Q", "3台/4台×22h×6天"),
    ("150T",      "X", "Y", "Z", "AA", "1台×22h×6/7天"),
    ("90T",       "AG", "AH", "AI", "AJ", "1台/2台×22h×6天"),
    ("150T-280T", "AU", "AV", "AW", "AX", "4台/5台×22h×6天"),
]
machine = {"weeks": [], "groups": {}}
for name, cR, cLo, cHi, cBud, tag in groups:
    machine["groups"][name] = {"label": name, "capTag": tag, "real": [], "capLo": [], "capHi": [], "budget": []}

for r in ws2.iter_rows(min_row=5, max_row=90, values_only=True):
    wk = r[1]  # B列 周
    if wk is None:
        continue
    wk = str(wk).strip()
    if not wk.startswith("202"):
        continue
    machine["weeks"].append(wk)
    for name, cR, cLo, cHi, cBud, tag in groups:
        g = machine["groups"][name]
        g["real"].append(num(r[openpyxl.utils.column_index_from_string(cR)-1]))
        g["capLo"].append(num(r[openpyxl.utils.column_index_from_string(cLo)-1]))
        g["capHi"].append(num(r[openpyxl.utils.column_index_from_string(cHi)-1]))
        g["budget"].append(num(r[openpyxl.utils.column_index_from_string(cBud)-1]))

# ============ 2. 原料KG（分析3） ============
# 注：旧版源文件有「分析3-品种分析原料分析」表；新版已并入「分析4-注塑半成品原料库存水位」，
#     网页原料需求表改由 EMBEDDED_GAP.material（_extract_gap.py）提供，此处保留空结构以兼容 data.json 字段。
material = {"weeks": [], "items": []}
if "分析3-品种分析原料分析" in wb.sheetnames:
    ws3 = wb["分析3-品种分析原料分析"]
    # 周表头在 R4, F列(6)起
    hdr = [c.value for c in ws3[4]]
    mat_weeks = []
    for c in range(6, 43+1):  # F..AQ
        v = hdr[c-1]
        if v is not None:
            mat_weeks.append(str(v).strip())
    material["weeks"] = mat_weeks
    for r in ws3.iter_rows(min_row=5, values_only=True):
        code = r[1]
        name = r[2]
        if code is None and name is None:
            continue
        vals = [num(r[c-1]) for c in range(6, 43+1)]
        material["items"].append({
            "code": str(code) if code is not None else "",
            "name": str(name) if name is not None else "",
            "usage": str(r[0]) if r[0] is not None else "",
            "desc": str(r[3]) if r[3] is not None else "",
            "total": num(r[4]),
            "weekly": vals,
        })
else:
    print("提示: 源文件无「分析3-品种分析原料分析」表，原料KG字段输出为空（网页已改用 EMBEDDED_GAP.material）")

# ============ 3. 产值（原始数据5） ============
ws5 = wb["原始数据5（外壳产值单价）"]
pv_month = defaultdict(float)       # 年月 -> 金额
pv_qty = defaultdict(float)
pv_cat = defaultdict(float)         # 品名 -> 金额
for r in ws5.iter_rows(min_row=2, values_only=True):
    dt = r[0]
    amt = num(r[14])  # O列 金额
    qty = num(r[9])   # J列 入库数量
    cat = str(r[6]).strip() if r[6] is not None else "其他"  # G列 品名
    if amt is None and qty is None:
        continue
    ym = None
    if isinstance(dt, datetime):
        ym = dt.strftime("%Y-%m")
    elif dt is not None:
        s = str(dt)
        m = re.match(r"(\d{4})-(\d{2})", s)
        if m:
            ym = f"{m.group(1)}-{m.group(2)}"
    if ym:
        pv_month[ym] += (amt or 0)
        pv_qty[ym] += (qty or 0)
    pv_cat[cat] += (amt or 0)

pv_month_sorted = sorted(pv_month.items())
productionValue = {
    "months": [k for k, v in pv_month_sorted],
    "amount": [round(v, 2) for k, v in pv_month_sorted],
    "qty": [round(pv_qty[k], 2) for k, v in pv_month_sorted],
    "byCategory": [{"name": k, "value": round(v, 2)} for k, v in sorted(pv_cat.items(), key=lambda x: -x[1])],
}

# ============ 4. 备料明细汇总 + 模具实单工时（原始数据1） ============
ws1 = wb["原始数据1（工单备料明细）"]
byCat = Counter()     # 品名 -> 扣库存后待生产数量
byMach = Counter()    # 机台 -> 行数
byWeek = Counter()    # 周 -> 扣库存后待生产数量
moldHours = defaultdict(lambda: defaultdict(float))  # 模具 -> 周 -> 工时
moldQty = defaultdict(lambda: defaultdict(float))
totalOrders = 0
totalBacklog = 0.0

def extract_mold(spec, pname):
    """从规格/品名提取模具编号，如 DH010S / DH010M / DH001"""
    if spec:
        s = str(spec)
        for m in re.findall(r"DH\d{3}[A-Z]?", s):
            return m
    if pname:
        p = str(pname)
        m = re.search(r"DH\d{3}[A-Z]?", p)
        if m:
            return m.group(0)
    return None

for r in ws1.iter_rows(min_row=4, values_only=True):
    c = r[2]  # 发料料号
    d = r[3]  # 品名
    if c is None and d is None:
        continue
    totalOrders += 1
    cat = str(d).strip() if d is not None else "其他"
    mach = str(r[16]).strip() if r[16] is not None else "0"
    wk = str(r[15]).strip() if r[15] is not None else ""
    l_val = num(r[11])  # 扣库存后待生产
    j_val = num(r[9])   # 订单生产工时
    spec = r[4]         # 规格
    byCat[cat] += (l_val or 0)
    byMach[mach] += 1
    if wk:
        byWeek[wk] += (l_val or 0)
    totalBacklog += (l_val or 0)
    mold = extract_mold(spec, d)
    if mold and wk:
        moldHours[mold][wk] += (j_val or 0)
        moldQty[mold][wk] += (l_val or 0)

all_weeks_mold = sorted(set(w for m in moldHours.values() for w in m.keys()))
moldList = []
for m in sorted(moldHours.keys()):
    moldList.append({
        "mold": m,
        "weeks": all_weeks_mold,
        "hours": [round(moldHours[m].get(w, 0), 2) for w in all_weeks_mold],
        "qty": [round(moldQty[m].get(w, 0), 1) for w in all_weeks_mold],
    })

backlog = {
    "totalOrders": totalOrders,
    "totalBacklog": round(totalBacklog, 0),
    "byCategory": [{"name": k, "value": round(v, 1)} for k, v in byCat.most_common()],
    "byMachine": [{"name": k, "value": v} for k, v in byMach.most_common()],
    "byWeek": [{"name": k, "value": round(v, 1)} for k, v in sorted(byWeek.items())],
}

# ============ 输出 ============
result = {
    "meta": {
        "generatedAt": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "source": "最新排程产能分析-公式模板-最新.xlsm",
        "totalOrders": totalOrders,
        "machineGroups": list(groups[i][0] for i in range(len(groups))),
    },
    "machine": machine,
    "material": material,
    "productionValue": productionValue,
    "mold": {"weeks": all_weeks_mold, "items": moldList},
    "backlog": backlog,
}

import os
os.makedirs(os.path.dirname(OUT), exist_ok=True)
with open(OUT, "w", encoding="utf-8") as f:
    json.dump(result, f, ensure_ascii=False)

print("输出:", OUT)
print("机台周数:", len(machine["weeks"]), machine["weeks"][:3], "...", machine["weeks"][-3:])
for name, g in machine["groups"].items():
    print(f"  {name}: real样例={g['real'][:5]} budget样例={g['budget'][:5]}")
print("原料数:", len(material["items"]), "周数:", len(material["weeks"]))
print("原料周样例:", material["weeks"][:3], "...", material["weeks"][-3:])
print("产值月数:", len(productionValue["months"]), productionValue["months"][:3], "...", productionValue["months"][-3:])
print("产值品类数:", len(productionValue["byCategory"]))
print("模具数:", len(moldList), "模具周数:", len(all_weeks_mold))
print("模具样例:", [m["mold"] for m in moldList[:20]])
print("备料品名:", [(k, v) for k, v in byCat.most_common()])
print("备料机台:", [(k, v) for k, v in byMach.most_common()])
print("订单数:", totalOrders, "扣库存后待生产总计:", round(totalBacklog, 0))
print("JSON大小(KB):", os.path.getsize(OUT)//1024)
wb.close()
