# -*- coding: utf-8 -*-
"""抽取 VLOOKUP 映射表（产能参数/上线日期/BOM），供网页复现 J-W 列公式"""
import openpyxl, json, os, re
from datetime import datetime, date

SRC = r"D:\WorkBuddy工作资料\injection_dashboard\最新排程产能分析-公式模板-最新.xlsm"
OUT = r"D:\WorkBuddy工作资料\injection_dashboard\mappings.json"

wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)

def s(v):
    if v is None: return None
    if isinstance(v, datetime): return v.strftime("%Y-%m-%d")
    if isinstance(v, date): return v.strftime("%Y-%m-%d")
    return str(v).strip()

def map_machine(t):
    """把原始数据4 H列的具体吨位映射回网页笼统分组（改善5）"""
    if t is None:
        return None
    m = str(t).strip().upper()
    mm = re.match(r'(\d+)\s*T', m)
    if not mm:
        return None
    tons = int(mm.group(1))
    if tons in (85, 90):
        return "90T"
    if tons == 150:
        return "150T"
    if tons in (230, 260):
        return "230T-280T"
    if tons in (400, 450, 480, 550):
        return "400T-550T"
    return m

# 1. 产能参数映射：料号A -> {品种E, 机台H, 穴数K, 周期M, 毛重O}
ws4 = wb["原始数据4(产能参数)"]
cap = {}
for r in ws4.iter_rows(min_row=2, values_only=True):
    code = s(r[0])
    if code is None: continue
    mach = map_machine(r[7])
    if not mach:  # 只保留注塑件（0/空/#N/A 映射为 None）
        continue
    if code not in cap:   # VLOOKUP 取第一个匹配，避免多行料号被后行覆盖
        cap[code] = {
            "p": s(r[4]),      # 品种(备注/简称)
            "m": mach,         # 机台（已映射到笼统分组）
            "c": s(r[10]),     # 穴数(如 1*1)
            "t": s(r[12]),     # 周期秒
            "w": s(r[14]),     # 毛重
        }

# 2. 上线日期映射：订单号A -> 上线日期B
ws2 = wb["原始数据2(最新排程上线日期)"]
line = {}
for r in ws2.iter_rows(min_row=2, values_only=True):
    key = s(r[0])
    if key is None: continue
    v = s(r[1])
    if v not in (None, ""):
        line[key] = v

# 3. BOM映射：主件C -> {元件D(原料代码), 元件规格K(原料名称)}，保留第一个匹配
wsb = wb["BOM计算"]
bom = {}
for r in wsb.iter_rows(min_row=3, values_only=True):
    key = s(r[2])
    if key is None: continue
    if key not in bom:
        bom[key] = {"c": s(r[3]), "n": s(r[10])}

# 4. 验证：用原始数据1 抽料号，比对映射结果 vs 缓存值
ws1 = wb["原始数据1（工单备料明细）"]
print("="*100)
print("验证映射正确性（抽查原始数据1 前6条有效料号）:")
checked = 0
for r in ws1.iter_rows(min_row=4, values_only=True):
    code = s(r[2])
    if code is None: continue
    if checked >= 6: break
    checked += 1
    cm = cap.get(code, {})
    print(f"  料号={code}")
    print(f"    原值: 品种={s(r[13])} 机台={s(r[16])} 穴数={s(r[17])} 周期={s(r[18])} 毛重={s(r[19])} 原料码={s(r[20])}")
    print(f"    映射: 品种={cm.get('p')} 机台={cm.get('m')} 穴数={cm.get('c')} 周期={cm.get('t')} 毛重={cm.get('w')} 原料码={bom.get(code,{}).get('c')}")

# 5. 输出
result = {"capacity": cap, "lineDate": line, "bom": bom}
os.makedirs(os.path.dirname(OUT), exist_ok=True)
with open(OUT, "w", encoding="utf-8") as f:
    json.dump(result, f, ensure_ascii=False, separators=(",", ":"))

print("\n输出:", OUT)
print("产能参数注塑件映射数:", len(cap))
print("上线日期映射数:", len(line))
print("BOM主件映射数:", len(bom))
print("文件大小(KB):", os.path.getsize(OUT)//1024)
wb.close()
